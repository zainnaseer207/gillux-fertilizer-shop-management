"""
utils/excel_exporter.py

Generic Excel export function — koi bhi report (headers + rows) isko
call karke .xlsx bana sakta hai.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def export_to_excel(headers: list, rows: list, output_path: str, sheet_title: str = "Report"):
    """
    headers: ["Column1", "Column2", ...]
    rows: [["val1", "val2", ...], ...]  (list of lists — string/number)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws.append(headers)
    header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    for row in rows:
        ws.append(row)

    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max(12, min(max_length + 2, 40))

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path